from __future__ import annotations

from pathlib import Path
import re
from typing import Iterable

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_DIR = Path("/Users/jack/Downloads/国开行投标简历")
OUTPUT_PATH = BASE_DIR / "国开行投标简历项目汇总.xlsx"


def clean(text: str) -> str:
    text = (text or "").replace("\xa0", " ").replace("\r", "\n")
    lines = [line.strip() for line in text.split("\n")]
    text = " ".join(line for line in lines if line)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def unique_nonempty(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        value = clean(value)
        if not value or value in seen:
            continue
        seen.add(value)
        result.append(value)
    return result


def row_texts(row) -> list[str]:
    return [clean(cell.text) for cell in row.cells]


def find_value_after_label(table, label: str) -> str:
    for row in table.rows:
        vals = row_texts(row)
        for idx, val in enumerate(vals):
            if val == label:
                for nxt in vals[idx + 1 :]:
                    if nxt and nxt != label:
                        return nxt
    return ""


def extract_name_from_filename(path: Path) -> str:
    name = path.stem
    name = name.replace("+new", "")
    patterns = [
        r"国家开发银行人员简历[_-]?(.+?)_(?:\d{8}|\d{6})(?:_.+)?$",
        r"国开行人员团队简历-\s*(.+)$",
        r"人员团队简历-(.+)$",
        r"\d+\+(.+?)\+工作简历",
    ]
    for pattern in patterns:
        m = re.search(pattern, name)
        if m:
            return m.group(1).strip()
    return name.strip()


def project_header_index(table) -> int | None:
    for idx, row in enumerate(table.rows):
        vals = row_texts(row)
        joined = " ".join(vals)
        if "项目名称及项目内容" in joined or ("日期" in joined and "项目" in joined):
            return idx
    return None


def extract_projects(table) -> list[str]:
    start = project_header_index(table)
    if start is None:
        return []

    projects: list[str] = []
    for row in table.rows[start + 1 :]:
        vals = row_texts(row)
        nonempty = unique_nonempty(vals)
        if not nonempty:
            continue

        # Common templates place project text in the 2nd logical column,
        # but merged cells duplicate values, so we choose the first strong
        # candidate after skipping date-like values and generic headers.
        candidate = ""
        for val in nonempty:
            if val in {
                "日期",
                "项目名称及项目内容",
                "项目名称",
                "项目所属机构名称",
                "担任何职",
                "岗位职责",
                "岗位工作时长（月）",
                "相关领域主要工作经历",
            }:
                continue
            if re.fullmatch(r"[\d./-]+(?:至今)?", val):
                continue
            if re.fullmatch(r"\d+", val):
                continue
            candidate = val
            break

        if not candidate:
            continue

        # Skip rows that are really role/institution duplicates.
        if candidate in {"系统开发", "开发工程师", "项目经理", "测试工程师", "需求分析师"}:
            continue

        projects.append(candidate)

    return unique_nonempty(projects)


def extract_resume(path: Path) -> tuple[str, str, list[str]]:
    doc = Document(str(path))
    if not doc.tables:
        return extract_name_from_filename(path), "", []

    table = doc.tables[0]
    name = find_value_after_label(table, "姓名") or extract_name_from_filename(path)
    role = find_value_after_label(table, "拟投入本项目工作岗位")
    if not role:
        role = find_value_after_label(table, "拟在本项目任职")
    projects = extract_projects(table)
    return clean(name), clean(role), projects


def build_workbook(rows: list[tuple[int, str, str, str, str]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "项目汇总"
    ws.append(["序号", "文件夹名称", "姓名", "拟投入本项目工作岗位", "项目名称及项目内容"])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append(list(row))

    widths = {
        "A": 8,
        "B": 28,
        "C": 18,
        "D": 28,
        "E": 72,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    return wb


def main() -> None:
    files = sorted(
        p
        for p in BASE_DIR.rglob("*.docx")
        if not p.name.startswith("~$")
        and not p.name.startswith(".~")
        and not p.name.startswith(".")
    )
    rows: list[tuple[int, str, str, str, str]] = []
    seq = 1
    for path in files:
        folder_name = path.parent.name
        name, role, projects = extract_resume(path)
        if projects:
            for project in projects:
                rows.append((seq, folder_name, name, role, project))
                seq += 1
        else:
            rows.append((seq, folder_name, name, role, ""))
            seq += 1

    wb = build_workbook(rows)
    wb.save(OUTPUT_PATH)
    print(f"files={len(files)}")
    print(f"rows={len(rows)}")
    print(f"output={OUTPUT_PATH}")


if __name__ == "__main__":
    main()
