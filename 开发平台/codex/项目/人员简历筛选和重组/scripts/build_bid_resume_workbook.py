from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


DATE_RE = re.compile(r"(\d{4})[./](\d{2})\s*[-—–~至到]+\s*(?:(\d{4})[./](\d{2})|(至今|现在))")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SKIP_NAMES = ("~$", ".~", ".")


def clean(text: str) -> str:
    text = (text or "").replace("\xa0", " ").replace("\t", " ").replace("\r", "\n")
    lines = [x.strip() for x in text.split("\n") if x.strip()]
    text = " ".join(lines)
    text = text.replace("，", ",").replace("；", ";")
    return re.sub(r"\s+", " ", text).strip(" ,;")


def is_real_docx(path: Path) -> bool:
    return path.suffix.lower() == ".docx" and not path.name.startswith(SKIP_NAMES)


def extract_name(path: Path, table_text: str = "") -> str:
    m = re.search(r"姓名\s*[：:]?\s*([\u4e00-\u9fffA-Za-z0-9]+)", table_text)
    if m:
        return m.group(1)
    stem = path.stem.replace("+new", "")
    for pat in [
        r"国家开发银行人员简历[_-]?(.+?)_(?:\d{8}|\d{6})(?:_.+)?$",
        r"国开行人员团队简历-\s*(.+)$",
        r"人员团队简历-(.+)$",
        r"\d+\+(.+?)\+工作简历",
    ]:
        mm = re.search(pat, stem)
        if mm:
            return mm.group(1).strip()
    return stem


def parse_period(text: str):
    text = clean(text).replace(" - ", "-").replace("—", "-").replace("–", "-")
    m = DATE_RE.search(text)
    if not m:
        return None
    sy, sm, ey, em, current = m.groups()
    start = int(sy) * 12 + int(sm)
    if current:
        end = 999999
        norm = f"{sy}/{sm}-至今"
    else:
        end = int(ey) * 12 + int(em)
        norm = f"{sy}/{sm}-{ey}/{em}"
    return {"norm": norm, "start": start, "end": end}


def row_texts(row) -> list[str]:
    return [clean(cell.text) for cell in row.cells]


def find_value_after_label(table, label: str) -> str:
    for row in table.rows:
        vals = row_texts(row)
        for idx, val in enumerate(vals):
            if val == label:
                for nxt in vals[idx + 1 :]:
                    if nxt and nxt != label:
                        return nxt
    return ""


def project_header_index(table) -> int | None:
    for idx, row in enumerate(table.rows):
        vals = row_texts(row)
        joined = " ".join(vals)
        if "项目名称及项目内容" in joined or ("日期" in joined and "项目" in joined):
            return idx
    return None


def extract_projects(table) -> list[dict]:
    start = project_header_index(table)
    if start is None:
        return []
    out = []
    seen = set()
    for row in table.rows[start + 1 :]:
        vals = row_texts(row)
        nonempty = [v for v in vals if v]
        if not nonempty:
            continue
        parsed = None
        project = ""
        for val in nonempty:
            if not parsed:
                parsed = parse_period(val)
                if parsed:
                    continue
            if not project and val not in {
                "日期",
                "项目名称及项目内容",
                "项目名称",
                "项目所属机构名称",
                "担任何职",
                "岗位职责",
                "岗位工作时长（月）",
                "相关领域主要工作经历",
            }:
                if not parse_period(val):
                    project = val
                    break
        if project and parsed:
            key = (parsed["norm"], project)
            if key not in seen:
                seen.add(key)
                out.append(
                    {
                        "period": parsed["norm"],
                        "start": parsed["start"],
                        "end": parsed["end"],
                        "project": project,
                    }
                )
    return out


def get_main_text(table) -> str:
    candidates = []
    for row in table.rows[:8]:
        vals = [cell.text for cell in row.cells]
        if any("主要经历专业技能" in val or "主要经历" in val for val in vals):
            for val in vals:
                cv = clean(val)
                if "主要经历" in cv and len(cv) > 20:
                    candidates.append(cv)
    if not candidates:
        return ""
    text = max(candidates, key=len)
    m = re.search(r"主要经历[：: ]*(.*?)(?:专业技能[：: ]*|主要技能[：: ]*|$)", text, re.S)
    return clean(m.group(1) if m else text)


def parse_work_history(main_text: str) -> list[dict]:
    if not main_text:
        return []
    text = main_text.replace(" - ", "-").replace("—", "-").replace("–", "-")
    parts = re.split(r"(?=\d{4}[./]\d{2}\s*[-~至到]+\s*(?:\d{4}[./]\d{2}|至今|现在))", text)
    out = []
    seen = set()
    for part in parts:
        part = clean(part)
        if not part:
            continue
        m = DATE_RE.search(part)
        if not m:
            continue
        period_obj = parse_period(m.group(0))
        if not period_obj:
            continue
        rest = clean(part[m.end() :]).lstrip(",")
        company = clean(rest.split(",")[0]) if rest else ""
        key = (period_obj["norm"], company)
        if key not in seen:
            seen.add(key)
            out.append(
                {
                    "period": period_obj["norm"],
                    "start": period_obj["start"],
                    "end": period_obj["end"],
                    "company": company,
                }
            )
    return out


def style_sheet(ws, widths: dict[str, int]) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def gather_resume_data(base_dir: Path):
    files = sorted(p for p in base_dir.rglob("*.docx") if is_real_docx(p))
    summary_rows = []
    people = defaultdict(lambda: {"companies": [], "projects": []})

    seq = 1
    for path in files:
        try:
            doc = Document(str(path))
        except Exception:
            continue
        if not doc.tables:
            continue
        table = doc.tables[0]
        head = "\n".join(cell.text for row in table.rows[:6] for cell in row.cells)
        name = clean(find_value_after_label(table, "姓名")) or extract_name(path, head)
        role = clean(find_value_after_label(table, "拟投入本项目工作岗位"))
        if not role:
            role = clean(find_value_after_label(table, "拟在本项目任职"))
        folder = path.parent.name
        projects = extract_projects(table)
        if projects:
            for project in projects:
                summary_rows.append((seq, folder, name, role, project["project"]))
                seq += 1
        else:
            summary_rows.append((seq, folder, name, role, ""))
            seq += 1
        people[name]["companies"].extend(parse_work_history(get_main_text(table)))
        people[name]["projects"].extend(projects)

    for name, data in people.items():
        comp_seen = set()
        comps = []
        for company in data["companies"]:
            key = (company["period"], company["company"])
            if key not in comp_seen:
                comp_seen.add(key)
                comps.append(company)
        proj_seen = set()
        projs = []
        for project in data["projects"]:
            key = (project["period"], project["project"])
            if key not in proj_seen:
                proj_seen.add(key)
                projs.append(project)
        data["companies"] = sorted(comps, key=lambda x: (x["start"], x["end"]), reverse=True)
        data["projects"] = sorted(projs, key=lambda x: (x["start"], x["end"]), reverse=True)

    return files, summary_rows, people


def build_workbook(base_dir: Path, output_path: Path) -> None:
    _, summary_rows, people = gather_resume_data(base_dir)
    wb = Workbook()

    # 项目汇总
    ws = wb.active
    ws.title = "项目汇总"
    ws.append(["序号", "文件夹名称", "姓名", "拟投入本项目工作岗位", "项目名称及项目内容"])
    for row in summary_rows:
        ws.append(list(row))
    style_sheet(ws, {"A": 8, "B": 28, "C": 18, "D": 28, "E": 72})

    # 项目统计
    stat = wb.create_sheet("项目统计")
    stat.append(["序号", "项目名称及项目内容", "人数", "人员名单"])
    project_people = defaultdict(list)
    seen = defaultdict(set)
    for _, _, name, _, project in summary_rows:
        if not project:
            continue
        if name not in seen[project]:
            seen[project].add(name)
            project_people[project].append(name)
    items = sorted(project_people.items(), key=lambda kv: (-len(kv[1]), kv[0]))
    for idx, (project, names) in enumerate(items, start=1):
        stat.append([idx, project, len(names), "、".join(names)])
    style_sheet(stat, {"A": 8, "B": 60, "C": 10, "D": 80})

    # 跨文件夹同岗位
    cross = wb.create_sheet("跨文件夹同岗位")
    cross.append(["序号", "姓名", "拟投入本项目工作岗位", "涉及文件夹数", "文件夹名称"])
    mapping = defaultdict(lambda: defaultdict(set))
    for _, folder, name, role, _ in summary_rows:
        if name and role and folder:
            mapping[name][role].add(folder)
    idx = 1
    for name, roles in mapping.items():
        for role, folders in roles.items():
            if len(folders) > 1:
                cross.append([idx, name, role, len(folders), "、".join(sorted(folders))])
                idx += 1
    style_sheet(cross, {"A": 8, "B": 16, "C": 28, "D": 12, "E": 80})

    # 项目岗位重复标注
    dup = wb.create_sheet("项目岗位重复标注")
    dup.append(
        [
            "序号",
            "文件夹名称",
            "姓名",
            "拟投入本项目工作岗位",
            "项目名称及项目内容",
            "是否项目岗位相同",
            "同组人数",
            "同组姓名",
        ]
    )
    groups = defaultdict(list)
    for _, folder, name, role, project in summary_rows:
        if project and role:
            groups[(project, role)].append((folder, name))
    for idx, (_, folder, name, role, project) in enumerate(summary_rows, start=1):
        group = groups.get((project, role), []) if project and role else []
        unique_names = []
        name_seen = set()
        for _, n in group:
            if n and n not in name_seen:
                name_seen.add(n)
                unique_names.append(n)
        dup.append(
            [
                idx,
                folder,
                name,
                role,
                project,
                "是" if len(group) > 1 else "否",
                len(group) if group else 0,
                "、".join(unique_names),
            ]
        )
    style_sheet(dup, {"A": 8, "B": 24, "C": 16, "D": 28, "E": 60, "F": 16, "G": 10, "H": 60})

    # 主要经历工作经历 + 未映射
    main = wb.create_sheet("主要经历工作经历")
    main.append(["序号", "人员姓名", "公司周期", "公司名称", "项目周期", "项目名称及项目内容"])
    unmatched_sheet = wb.create_sheet("项目周期未映射")
    unmatched_sheet.append(["序号", "人员姓名", "项目周期", "项目名称及项目内容"])
    row_idx = 1
    unmatched_idx = 1
    for name in sorted(people):
        companies = people[name]["companies"]
        projects = people[name]["projects"]
        if not companies:
            for project in projects:
                main.append([row_idx, name, "", "", project["period"], project["project"]])
                row_idx += 1
            continue
        matched = set()
        for company in companies:
            company_matches = []
            for project in projects:
                if project["start"] >= company["start"] and project["end"] <= company["end"]:
                    company_matches.append(project)
                    matched.add((project["period"], project["project"]))
            if company_matches:
                for project in company_matches:
                    main.append(
                        [
                            row_idx,
                            name,
                            company["period"],
                            company["company"],
                            project["period"],
                            project["project"],
                        ]
                    )
                    row_idx += 1
            else:
                main.append([row_idx, name, company["period"], company["company"], "", ""])
                row_idx += 1
        for project in projects:
            key = (project["period"], project["project"])
            if key not in matched:
                unmatched_sheet.append([unmatched_idx, name, project["period"], project["project"]])
                unmatched_idx += 1
    style_sheet(main, {"A": 8, "B": 16, "C": 20, "D": 34, "E": 20, "F": 60})
    style_sheet(unmatched_sheet, {"A": 8, "B": 16, "C": 20, "D": 60})

    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("base_dir", help="Resume folder root")
    parser.add_argument(
        "--output",
        help="Output workbook path; defaults to <base_dir>/国开行投标简历项目汇总.xlsx",
        default=None,
    )
    args = parser.parse_args()

    base_dir = Path(args.base_dir).expanduser()
    output = Path(args.output).expanduser() if args.output else base_dir / "国开行投标简历项目汇总.xlsx"
    build_workbook(base_dir, output)
    print(f"output={output}")


if __name__ == "__main__":
    main()
