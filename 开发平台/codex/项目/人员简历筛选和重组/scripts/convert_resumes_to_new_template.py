from __future__ import annotations

import argparse
import copy
import math
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


TODAY = date(2026, 3, 25)
FONT_NAME = "仿宋"
FONT_SIZE = Pt(12)


@dataclass
class CertRow:
    name: str
    level: str
    major: str
    note: str


@dataclass
class ProjectRow:
    date_text: str
    name: str
    institution: str
    role: str
    duty: str
    months: str
    content: str


@dataclass
class ResumeData:
    name: str
    age: str
    id_number: str
    school: str
    degree: str
    work_start: str
    work_years: str
    related_years: str
    target_role: str
    role_years: str
    summary: str
    cert_summary: str
    certs: list[CertRow]
    projects: list[ProjectRow]


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\n", " ").strip())


def month_diff(start_year: int, start_month: int, end_year: int, end_month: int) -> int:
    return max(1, (end_year - start_year) * 12 + (end_month - start_month) + 1)


def compute_age_from_birth(year: int, month: int, day: int) -> str:
    years = TODAY.year - year
    if (TODAY.month, TODAY.day) < (month, day):
        years -= 1
    return str(max(0, years))


def compute_age_from_id(id_number: str) -> str:
    value = re.sub(r"\s+", "", id_number)
    if len(value) == 18 and value[:17].isdigit():
        return compute_age_from_birth(int(value[6:10]), int(value[10:12]), int(value[12:14]))
    if len(value) == 15 and value[:15].isdigit():
        year = 1900 + int(value[6:8])
        month = int(value[8:10])
        day = int(value[10:12])
        return compute_age_from_birth(year, month, day)
    return ""


def normalize_school_text(text: str) -> str:
    value = normalize_text(text)
    value = re.sub(r"^\d{4}年毕业于", "", value)
    value = re.sub(r"^\d{4}[.\-]\d{1,2}年?毕业于", "", value)
    value = value.replace("学校", "").strip()
    value = re.sub(r"\s+", " ", value)
    tokens = [t for t in value.split(" ") if t]
    if not tokens:
        return value
    school = tokens[0]
    major = "".join(tokens[1:]).replace("专业专业", "专业")
    if major and not major.endswith("专业"):
        major = f"{major}专业"
    if major:
        return f"{school}，{major}"
    return school


def normalize_degree_text(text: str) -> str:
    value = normalize_text(text)
    value = value.replace("、", "，").replace(",", "，")
    value = re.sub(r"\s+", "", value)
    pairs = [
        ("本科学历", "学士学位"),
        ("硕士学历", "硕士学位"),
        ("博士学历", "博士学位"),
        ("专科学历", ""),
    ]
    for edu, degree in pairs:
        if edu in value:
            if degree:
                return f"{edu}，{degree}"
            return edu
    return value


def parse_date_range(text: str) -> tuple[int, int, int, int]:
    text = text.strip().replace("—", "-").replace("–", "-").replace("至今", f"{TODAY.year}.{TODAY.month:02d}")
    text = text.replace("年", ".").replace("月", "").replace(" ", "")
    m = re.match(r"(\d{4})\.(\d{1,2})-(\d{4})\.(\d{1,2})", text)
    if not m:
        raise ValueError(f"Unrecognized date range: {text}")
    return tuple(int(x) for x in m.groups())  # type: ignore[return-value]


def extract_institution(project_name: str) -> str:
    name = normalize_text(project_name)
    patterns = [
        r"^(.+?银行)",
        r"^(.+?农商行)",
        r"^(.+?农商银行)",
        r"^(.+?保险)",
        r"^(.+?证券)",
        r"^(.+?金租)",
        r"^(.+?集团)",
        r"^(.+?公司)",
    ]
    for pattern in patterns:
        m = re.match(pattern, name)
        if m:
            return m.group(1)
    return name[:20]


def compress_content(text: str, limit: int = 28) -> str:
    text = normalize_text(text)
    text = re.sub(r"[；;，,。]+", "、", text)
    parts = [p for p in text.split("、") if p]
    if not parts:
        return text[:limit]
    brief = "、".join(parts[:3])
    return brief[:limit]


def compress_duty(text: str, limit: int = 42) -> str:
    text = normalize_text(text)
    text = re.sub(r"[；;，,。]+", "、", text)
    parts = [p for p in text.split("、") if p]
    if not parts:
        return text[:limit]
    duty = "、".join(parts[:4])
    return duty[:limit]


def compute_work_years(work_start: str) -> str:
    m = re.match(r"(\d{4})\.(\d{1,2})", work_start)
    if not m:
        return ""
    sy, sm = int(m.group(1)), int(m.group(2))
    months = month_diff(sy, sm, TODAY.year, TODAY.month)
    return str(months // 12)


def keyword_summary(projects: list[ProjectRow]) -> str:
    joined = " ".join([p.name + " " + p.duty + " " + p.content for p in projects[:8]])
    skills: list[str] = []
    if any(k in joined for k in ["风险", "RWA", "资本", "限额", "并表", "压力测试", "模型"]):
        skills.append("长期从事银行风险管理相关项目实施")
    if any(k in joined for k in ["需求", "分析", "建模"]):
        skills.append("具备需求分析与业务建模经验")
    if any(k in joined for k in ["开发", "ETL", "Java", "数据"]):
        skills.append("具备系统开发与数据处理能力")
    if any(k in joined for k in ["测试", "投产", "上线", "联调"]):
        skills.append("具备联调测试与投产支持经验")
    if any(k in joined for k in ["项目经理", "组长", "负责人", "管理"]):
        skills.append("能够承担项目协调与实施推进工作")
    if not skills:
        skills.append("具备银行信息系统实施与项目交付经验")
    return "；".join(skills) + "。"


def extract_resume_data(path: Path) -> ResumeData:
    doc = Document(str(path))
    table = doc.tables[0]

    name = normalize_text(table.rows[0].cells[1].text)
    id_number = normalize_text(table.rows[2].cells[1].text)
    age = compute_age_from_id(id_number)
    if not age:
        raw_age = normalize_text(table.rows[0].cells[3].text)
        m = re.match(r"(\d{4})[.\-](\d{1,2})[.\-](\d{1,2})", raw_age)
        if m:
            age = compute_age_from_birth(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        elif raw_age.isdigit():
            age = raw_age
        else:
            age = ""
    degree = normalize_degree_text(table.rows[0].cells[6].text)
    work_start = normalize_text(table.rows[1].cells[3].text)
    related_years = normalize_text(table.rows[1].cells[6].text)
    school = normalize_school_text(table.rows[4].cells[1].text)
    target_role = normalize_text(table.rows[5].cells[4].text)
    work_years = compute_work_years(work_start)
    role_years = work_years

    certs: list[CertRow] = []
    row_idx = 8
    while row_idx < len(table.rows):
        first = normalize_text(table.rows[row_idx].cells[0].text)
        if first == "相关领域主要工作经历":
            break
        if first and first != "证书名称":
            certs.append(
                CertRow(
                    name=first,
                    level=normalize_text(table.rows[row_idx].cells[1].text),
                    major=normalize_text(table.rows[row_idx].cells[3].text),
                    note=normalize_text(table.rows[row_idx].cells[5].text),
                )
            )
        row_idx += 1

    while row_idx < len(table.rows) and normalize_text(table.rows[row_idx].cells[0].text) != "日期":
        row_idx += 1
    row_idx += 1

    projects: list[ProjectRow] = []
    for i in range(row_idx, len(table.rows)):
        row = table.rows[i]
        date_text = normalize_text(row.cells[0].text)
        if not date_text:
            continue
        project_name = normalize_text(row.cells[1].text)
        role = normalize_text(row.cells[3].text)
        note = normalize_text(row.cells[5].text)
        if not project_name:
            continue
        try:
            sy, sm, ey, em = parse_date_range(date_text)
            months = str(month_diff(sy, sm, ey, em))
        except Exception:
            months = ""
        projects.append(
            ProjectRow(
                date_text=date_text,
                name=project_name,
                institution=extract_institution(project_name),
                role=role,
                duty=compress_duty(note),
                months=months,
                content=compress_content(note),
            )
        )

    cert_summary = "；".join(
        [f"{c.name}{('（' + c.level + '）') if c.level else ''}" for c in certs if c.name]
    )
    if not cert_summary:
        cert_summary = "无"

    summary = keyword_summary(projects)
    return ResumeData(
        name=name,
        age=age,
        id_number=id_number,
        school=school,
        degree=degree,
        work_start=work_start,
        work_years=work_years,
        related_years=related_years,
        target_role=target_role,
        role_years=role_years,
        summary=summary,
        cert_summary=cert_summary,
        certs=certs or [CertRow(name="无", level="", major="", note="")],
        projects=projects,
    )


def ensure_run_font(paragraph) -> None:
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        run.font.name = FONT_NAME
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        run.font.size = FONT_SIZE


def set_cell_text(cell, text: str, align=WD_PARAGRAPH_ALIGNMENT.LEFT) -> None:
    cell.text = text
    for p in cell.paragraphs:
        p.alignment = align
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        ensure_run_font(p)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def clone_row_after(table, row_idx: int):
    template_row = table.rows[row_idx]
    new_tr = copy.deepcopy(template_row._tr)
    template_row._tr.addnext(new_tr)
    return table.rows[row_idx + 1]


def clear_row(row) -> None:
    seen = set()
    for cell in row.cells:
        tc_id = id(cell._tc)
        if tc_id in seen:
            continue
        seen.add(tc_id)
        set_cell_text(cell, "")


def fill_new_template(template_path: Path, data: ResumeData, output_path: Path) -> None:
    doc = Document(str(template_path))
    table = doc.tables[0]

    set_cell_text(table.rows[0].cells[1], data.name)
    set_cell_text(table.rows[0].cells[3], data.age)
    set_cell_text(table.rows[0].cells[7], data.id_number)

    set_cell_text(table.rows[1].cells[1], data.school)
    set_cell_text(table.rows[1].cells[7], data.degree)

    set_cell_text(table.rows[2].cells[1], data.work_years)
    set_cell_text(table.rows[2].cells[7], data.related_years)

    set_cell_text(table.rows[3].cells[1], data.target_role)
    set_cell_text(table.rows[3].cells[7], data.role_years)

    set_cell_text(table.rows[4].cells[1], data.summary)
    set_cell_text(table.rows[5].cells[1], data.cert_summary)

    base_project_row_index = 8
    project_rows_needed = max(3, len(data.projects))
    while len(table.rows) < 8 + project_rows_needed:
        clone_row_after(table, len(table.rows) - 1)

    for idx in range(8, len(table.rows)):
        clear_row(table.rows[idx])

    for i, project in enumerate(data.projects):
        row = table.rows[8 + i]
        set_cell_text(row.cells[0], project.date_text)
        set_cell_text(row.cells[1], project.name)
        set_cell_text(row.cells[3], project.institution)
        set_cell_text(row.cells[4], project.role)
        set_cell_text(row.cells[6], project.duty)
        set_cell_text(row.cells[7], project.months, WD_PARAGRAPH_ALIGNMENT.CENTER)

    # Keep the second table as a single-line roster for downstream paste/use.
    roster = doc.tables[1]
    if len(roster.rows) > 1:
        set_cell_text(roster.rows[1].cells[1], data.name)
        set_cell_text(roster.rows[1].cells[3], data.age)
        set_cell_text(roster.rows[1].cells[4], data.school)
        set_cell_text(roster.rows[1].cells[6], data.target_role)
        set_cell_text(roster.rows[1].cells[7], data.role_years)
        set_cell_text(roster.rows[1].cells[8], data.related_years)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def select_latest_files(source_root: Path) -> list[Path]:
    chosen: dict[tuple[str, str], tuple[float, Path]] = {}
    pattern = re.compile(r"国开行人员团队简历-(.+?)(?:_v(\d+(?:\.\d+)?))?\.docx$")
    for path in source_root.rglob("*.docx"):
        if "+new" in path.stem:
            continue
        m = pattern.match(path.name)
        if not m:
            continue
        person = m.group(1)
        version = float(m.group(2)) if m.group(2) else -1.0
        key = (str(path.parent.relative_to(source_root)), person)
        current = chosen.get(key)
        if current is None or version > current[0]:
            chosen[key] = (version, path)
    return [item[1] for item in sorted(chosen.values(), key=lambda x: str(x[1]))]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--source-root", required=True)
    parser.add_argument("--output-root", required=True)
    parser.add_argument("--suffix", default="")
    parser.add_argument("--flat", action="store_true")
    args = parser.parse_args()

    template_path = Path(args.template)
    source_root = Path(args.source_root)
    output_root = Path(args.output_root)

    processed = 0
    files = select_latest_files(source_root)
    for src in files:
        data = extract_resume_data(src)
        rel_dir = Path() if args.flat else src.parent.relative_to(source_root)
        stem = src.stem + args.suffix
        dst = output_root / rel_dir / f"{stem}{src.suffix}"
        fill_new_template(template_path, data, dst)
        processed += 1
        print(f"OK {src} -> {dst}")

    print(f"processed={processed}")


if __name__ == "__main__":
    main()
