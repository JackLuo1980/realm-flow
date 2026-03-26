from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt


FONT_NAME = "宋体"
FONT_SIZE = Pt(10.5)


@dataclass
class ImplProject:
    date_text: str
    name: str
    institution: str
    role: str
    duty: str
    months: str


@dataclass
class ImplResume:
    name: str
    age: str
    id_number: str
    school_major: str
    degree: str
    work_years: str
    related_years: str
    target_role: str
    role_years: str
    main_history: str
    skills_and_qualifications: str
    projects: list[ImplProject]


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\n", " ").replace("\t", " ").strip())


def month_diff(start_year: int, start_month: int, end_year: int, end_month: int) -> int:
    return max(1, (end_year - start_year) * 12 + (end_month - start_month) + 1)


def parse_period(text: str) -> tuple[int, int]:
    m = re.search(r"(\d{4})[./年-](\d{1,2})", text)
    if not m:
        raise ValueError(text)
    return int(m.group(1)), int(m.group(2))


def months_from_range(start: str, end: str) -> str:
    try:
        sy, sm = parse_period(start)
        if "至今" in end:
            ey, em = sy, sm
        else:
            ey, em = parse_period(end)
        return str(month_diff(sy, sm, ey, em))
    except Exception:
        return ""


def format_degree(value: str) -> str:
    text = normalize_text(value)
    mapping = {
        "本科": "本科学历、学士学位",
        "硕士": "硕士学历、硕士学位",
        "博士": "博士学历、博士学位",
        "专科": "专科",
    }
    return mapping.get(text, text)


def extract_institution(project_name: str) -> str:
    name = normalize_text(project_name)
    for pattern in [
        r"^(.+?银行)",
        r"^(.+?农商银行)",
        r"^(.+?农商行)",
        r"^(.+?保险)",
        r"^(.+?证券)",
        r"^(.+?集团)",
        r"^(.+?公司)",
    ]:
        m = re.match(pattern, name)
        if m:
            return m.group(1)
    return ""


def set_run_font(run) -> None:
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = FONT_SIZE


def set_cell_text(cell, text: str, align=WD_PARAGRAPH_ALIGNMENT.LEFT) -> None:
    cell.text = text
    for p in cell.paragraphs:
        p.alignment = align
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        if not p.runs:
            p.add_run("")
        for run in p.runs:
            set_run_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def parse_raw_resume(path: Path) -> ImplResume:
    doc = Document(str(path))
    table = doc.tables[0]

    name = normalize_text(table.rows[1].cells[1].text)
    work_years = normalize_text(table.rows[1].cells[4].text).replace("年", "")
    grad_time = normalize_text(table.rows[2].cells[1].text)
    school = normalize_text(table.rows[2].cells[4].text)
    major = normalize_text(table.rows[3].cells[1].text)
    degree = format_degree(normalize_text(table.rows[3].cells[4].text))
    title = normalize_text(table.rows[4].cells[4].text)

    school_major = school
    if major:
        school_major = f"{school}，{major}专业" if not major.endswith("专业") else f"{school}，{major}"

    work_history_rows: list[str] = []
    row_idx = 8
    while row_idx < len(table.rows):
        first = normalize_text(table.rows[row_idx].cells[0].text)
        if first == "项目经历（由近至远）":
            break
        if first and first != "开始时间":
            start = normalize_text(table.rows[row_idx].cells[0].text)
            end = normalize_text(table.rows[row_idx].cells[1].text)
            company = normalize_text(table.rows[row_idx].cells[2].text)
            role = normalize_text(table.rows[row_idx].cells[3].text)
            if company:
                work_history_rows.append(f"{start}-{end}，{company}，{role}".rstrip("，"))
        row_idx += 1
    main_history = " / ".join(work_history_rows)

    while row_idx < len(table.rows) and normalize_text(table.rows[row_idx].cells[0].text) != "开始时间":
        row_idx += 1
    row_idx += 1

    projects: list[ImplProject] = []
    while row_idx < len(table.rows):
        first = normalize_text(table.rows[row_idx].cells[0].text)
        if not first or first in {"能力与资质", "业务与技术能力详述", "资质认证", "参与培训", "技能标签"}:
            break
        start = normalize_text(table.rows[row_idx].cells[0].text)
        end = normalize_text(table.rows[row_idx].cells[1].text)
        name_text = normalize_text(table.rows[row_idx].cells[2].text)
        role = normalize_text(table.rows[row_idx].cells[3].text)
        duty = normalize_text(table.rows[row_idx].cells[4].text)
        if name_text:
            date_text = f"{start.replace('/', '.').replace('年', '.').replace('月', '')}-{'至今' if end == '至今' else end.replace('/', '.').replace('年', '.').replace('月', '')}"
            projects.append(
                ImplProject(
                    date_text=date_text,
                    name=name_text,
                    institution=extract_institution(name_text),
                    role=role,
                    duty=duty,
                    months=months_from_range(start, end),
                )
            )
        row_idx += 1

    skill_parts: list[str] = []
    while row_idx < len(table.rows):
        first = normalize_text(table.rows[row_idx].cells[0].text)
        if first in {"业务与技术能力详述", "资质认证", "参与培训", "技能标签"}:
            value = normalize_text(table.rows[row_idx].cells[1].text)
            if value:
                skill_parts.append(value)
        row_idx += 1

    return ImplResume(
        name=name,
        age="",
        id_number="",
        school_major=school_major,
        degree=degree,
        work_years=work_years,
        related_years=work_years,
        target_role=title,
        role_years=work_years,
        main_history=main_history,
        skills_and_qualifications=" / ".join(skill_parts),
        projects=projects,
    )


def fill_impl_template(template_path: Path, data: ImplResume, output_path: Path) -> None:
    doc = Document(str(template_path))
    table = doc.tables[0]

    set_cell_text(table.rows[0].cells[1], data.name)
    set_cell_text(table.rows[0].cells[3], data.age, WD_PARAGRAPH_ALIGNMENT.CENTER)
    set_cell_text(table.rows[0].cells[7], data.id_number)

    set_cell_text(table.rows[1].cells[1], data.school_major)
    set_cell_text(table.rows[1].cells[7], data.degree)

    set_cell_text(table.rows[2].cells[1], data.work_years)
    set_cell_text(table.rows[2].cells[7], data.related_years)

    set_cell_text(table.rows[3].cells[1], data.target_role)
    set_cell_text(table.rows[3].cells[7], data.role_years)

    set_cell_text(table.rows[4].cells[1], data.main_history)
    set_cell_text(table.rows[5].cells[1], data.skills_and_qualifications)

    for idx in range(8, len(table.rows)):
        for cell in table.rows[idx].cells:
            set_cell_text(cell, "")

    for idx, project in enumerate(data.projects[: len(table.rows) - 8]):
        row = table.rows[8 + idx]
        set_cell_text(row.cells[0], project.date_text)
        set_cell_text(row.cells[1], project.name)
        set_cell_text(row.cells[3], project.institution)
        set_cell_text(row.cells[4], project.role)
        set_cell_text(row.cells[6], project.duty)
        set_cell_text(row.cells[7], project.months, WD_PARAGRAPH_ALIGNMENT.CENTER)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("files", nargs="+")
    args = parser.parse_args()

    template = Path(args.template)
    for raw in [Path(f) for f in args.files]:
        data = parse_raw_resume(raw)
        output = raw.with_name(raw.stem + "+new.docx")
        fill_impl_template(template, data, output)
        print(f"OK {raw} -> {output}")


if __name__ == "__main__":
    main()
