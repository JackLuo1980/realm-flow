#!/usr/bin/env python3
import shutil
from pathlib import Path
from docx import Document

BASE_DIR = Path.home() / "Desktop" / "新版简历"
TEMPLATE_PATH = BASE_DIR / "简历模版.docx"
RESUME_1_0_DIR = BASE_DIR / "简历1.0"
ORIGINAL_RESUMES_DIR = BASE_DIR / "原始简历"
OUTPUT_DIR = BASE_DIR / "新简历"

RESUME_MAPPINGS = {
    "国开行人员团队简历-何文进.docx": "人员简历_何文进_202603.docx",
    "国开行人员团队简历-唐文波.docx": "02588+唐文波+工作简历.docx",
    "国开行人员团队简历-姜迪元.docx": "11919+姜迪元+工作简历.docx",
    "国开行人员团队简历-李倩倩.docx": "30280+李倩倩1+工作简历.docx",
    "国开行人员团队简历-李冬.docx": "02170+李冬+工作简历.docx",
    "国开行人员团队简历-李燕云.docx": "03017+李燕云+工作简历..docx",
    "国开行人员团队简历-荆芳.docx": "11237+荆芳+工作简历.docx",
}

def smart_extract_info(source_doc):
    info = {
        "name": "",
        "work_years": "",
        "school": "",
        "major": "",
        "title": "",
        "summary": "",
        "department": "",
        "graduation_date": ""
    }
    
    table = source_doc.tables[0]
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        
        for i in range(len(cells) - 1):
            cell_text = cells[i]
            next_cell_text = cells[i + 1]
            
            if ("姓" in cell_text and "名" in cell_text) and "工作年限" not in cell_text:
                if not info["name"] and next_cell_text and "年" not in next_cell_text:
                    info["name"] = next_cell_text
            
            elif "工作年限" in cell_text:
                if "年" in next_cell_text and not info["work_years"]:
                    info["work_years"] = next_cell_text
            
            elif "毕业学校" in cell_text and "学院" in next_cell_text or "大学" in next_cell_text:
                if not info["school"]:
                    info["school"] = next_cell_text
            
            elif ("专" in cell_text and "业" in cell_text) and cell_text != "专业":
                if not info["major"] and any(keyword in next_cell_text for keyword in ["工程", "技术", "管理"]):
                    info["major"] = next_cell_text
            
            elif "毕业时间" in cell_text:
                if "年" in next_cell_text or "月" in next_cell_text:
                    if not info["graduation_date"]:
                        info["graduation_date"] = next_cell_text
            
            elif "所在部门" in cell_text:
                if not info["department"] and "部" in next_cell_text:
                    info["department"] = next_cell_text
            
            elif "个人简介" in cell_text:
                if not info["summary"] and len(next_cell_text) > 10:
                    info["summary"] = next_cell_text
            
            elif "职称" in cell_text and "高级" in next_cell_text or "工程师" in next_cell_text:
                if not info["title"]:
                    info["title"] = next_cell_text
    
    return info

def create_resume_from_template(template_path, source_resume_path, output_path):
    shutil.copy2(template_path, output_path)
    template = Document(output_path)
    source_doc = Document(source_resume_path)
    
    personal_info = smart_extract_info(source_doc)
    
    print(f"提取的信息: {personal_info}")
    
    for i, table in enumerate(template.tables):
        for row_idx, row in enumerate(table.rows):
            for cell_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                
                if cell_text == "姓名" and personal_info["name"]:
                    cell.text = personal_info["name"]
                    print(f"填充姓名: {personal_info['name']}")
                
                elif cell_text == "毕业学校和专业" and personal_info["school"] and personal_info["major"]:
                    school_major = f"{personal_info['school']} {personal_info['major']}"
                    cell.text = school_major
                    print(f"填充毕业学校和专业: {school_major}")
                
                elif cell_text == "工作年限" and personal_info["work_years"]:
                    cell.text = personal_info["work_years"]
                    print(f"填充工作年限: {personal_info['work_years']}")
                
                elif cell_text == "主要经历专业技能" and personal_info["summary"]:
                    cell.text = personal_info["summary"]
                    print(f"填充主要经历专业技能: {personal_info['summary'][:50]}...")
    
    template.save(output_path)
    print(f"已生成: {output_path.name}")

def process_resumes():
    OUTPUT_DIR.mkdir(exist_ok=True)
    
    for resume_file, original_file in RESUME_MAPPINGS.items():
        resume_path = RESUME_1_0_DIR / resume_file
        original_path = ORIGINAL_RESUMES_DIR / original_file
        
        if not resume_path.exists():
            print(f"警告: 简历文件不存在: {resume_file}")
            continue
            
        if not original_path.exists():
            print(f"警告: 原始简历不存在: {original_file}, 使用当前简历内容")
            original_path = resume_path
        
        output_path = OUTPUT_DIR / resume_file
        create_resume_from_template(TEMPLATE_PATH, original_path, output_path)

if __name__ == "__main__":
    process_resumes()
    print(f"\n所有简历处理完成！输出目录: {OUTPUT_DIR}")
