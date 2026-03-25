#!/usr/bin/env python3
import os
import shutil
from pathlib import Path
from docx import Document

BASE_DIR = Path.home() / "Desktop" / "新版简历"
TEMPLATE_PATH = BASE_DIR / "简历模版.docx"
RESUME_1_0_DIR = BASE_DIR / "简历1.0"
ORIGINAL_RESUMES_DIR = BASE_DIR / "原始简历"
OUTPUT_DIR = BASE_DIR / "新简历"

RESUME_MAPPINGS = {
    "国开行人员团队简历-何文进.docx": "人员简历_何文进_202603.docx",
    "国开行人员团队简历-唐文波.docx": "02588+唐文波+工作简历.docx",
    "国开行人员团队简历-姜迪元.docx": "11919+姜迪元+工作简历.docx",
    "国开行人员团队简历-李倩倩.docx": "30280+李倩倩1+工作简历.docx",
    "国开行人员团队简历-李冬.docx": "02170+李冬+工作简历.docx",
    "国开行人员团队简历-李燕云.docx": "03017+李燕云+工作简历..docx",
    "国开行人员团队简历-荆芳.docx": "11237+荆芳+工作简历.docx",
}

def extract_table_data(doc):
    tables_data = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            table_data.append(row_data)
        tables_data.append(table_data)
    return tables_data

def create_resume_from_template(template_path, source_resume_path, output_path):
    template = Document(template_path)
    source_doc = Document(source_resume_path)
    
    source_tables = extract_table_data(source_doc)
    template_tables = extract_table_data(template)
    
    def find_value_by_key(source_tables, key):
        for source_table in source_tables:
            for source_row in source_table:
                for i, source_cell in enumerate(source_row):
                    if key.strip() in source_cell.strip():
                        if i + 1 < len(source_row):
                            value = source_row[i + 1].strip()
                            if value and value not in [key, source_cell]:
                                return value
                        return None
        return None
    
    if len(template_tables) >= 1 and len(source_tables) >= 1:
        for i, template_row in enumerate(template_tables[0]):
            for j, template_cell in enumerate(template_row):
                if template_cell.startswith("[") and template_cell.endswith("]"):
                    key = template_cell[1:-1]
                    found_value = find_value_by_key(source_tables, key)
                    
                    if found_value:
                        template.tables[0].rows[i].cells[j].text = found_value
    
    template.save(output_path)
    print(f"已生成: {output_path.name}")

def process_resumes():
    OUTPUT_DIR.mkdir(exist_ok=True)
    
    for resume_file, original_file in RESUME_MAPPINGS.items():
        resume_path = RESUME_1_0_DIR / resume_file
        original_path = ORIGINAL_RESUMES_DIR / original_file
        
        if not resume_path.exists():
            print(f"警告: 简历文件不存在: {resume_file}")
            continue
            
        if not original_path.exists():
            print(f"警告: 原始简历不存在: {original_file}, 使用当前简历内容")
            original_path = resume_path
        
        output_path = OUTPUT_DIR / resume_file
        create_resume_from_template(TEMPLATE_PATH, original_path, output_path)

if __name__ == "__main__":
    process_resumes()
    print(f"\n所有简历处理完成！输出目录: {OUTPUT_DIR}")
