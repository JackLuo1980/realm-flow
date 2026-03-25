#!/usr/bin/env python3
from docx import Document
from pathlib import Path

BASE_DIR = Path.home() / "Desktop" / "新版简历"
original_path = BASE_DIR / "原始简历" / "人员简历_何文进_202603.docx"

doc = Document(original_path)

print('=== 详细检查单元格内容 ===')
row = doc.tables[0].rows[1]
print(f'第1行单元格数量: {len(row.cells)}')

for i, cell in enumerate(row.cells):
    text = cell.text
    print(f'单元格{i}: 长度={len(text)}, 文本="{text}", repr={repr(text)}, 包含姓名={"姓名" in text}')
    
print('\n=== 检查所有行 ===')
for i, row in enumerate(doc.tables[0].rows[:5]):
    print(f'\n行{i}:')
    for j, cell in enumerate(row.cells[:5]):
        text = cell.text.strip()
        print(f'  列{j}: "{text}" (包含姓名: {"姓名" in text})')
