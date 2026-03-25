#!/usr/bin/env python3
import shutil
from pathlib import Path
from docx import Document

BASE_DIR = Path.home() / "Desktop" / "新版简历"
TEMPLATE_PATH = BASE_DIR / "简历模版.docx"
RESUME_1_0_DIR = BASE_DIR / "简历1.0"
ORIGINAL_RESUMES_DIR = BASE_DIR / "原始简历"
OUTPUT_DIR = BASE_DIR / "新简历"

RESUME_MAPPINGS = {
    "国开行人员团队简历-何文进.docx": "人员简历_何文进_202603.docx",
    "国开行人员团队简历-唐文波.docx": "02588+唐文波+工作简历.docx",
    "国开行人员团队简历-姜迪元.docx": "11919+姜迪元+工作简历.docx",
    "国开行人员团队简历-李倩倩.docx": "30280+李倩倩1+工作简历.docx",
    "国开行人员团队简历-李冬.docx": "02170+李冬+工作简历.docx",
    "国开行人员团队简历-李燕云.docx": "03017+李燕云+工作简历..docx",
    "国开行人员团队简历-荆芳.docx": "11237+荆芳+工作简历.docx",
}

def extract_personal_info(source_doc):
    info = {}
    for table in source_doc.tables:
        for row in table.rows:
            for i in range(len(row.cells) - 1):
                cell_text = row.cells[i].text.strip()
                next_cell_text = row.cells[i + 1].text.strip()
                
                if ("姓" in cell_text and "名" in cell_text) and cell_text not in ["姓名", "姓名" * len(cell_text)]:
                    if "name" not in info:
                        info["name"] = next_cell_text
                elif "工作年限" in cell_text and cell_text != "工作年限":
                    if "work_years" not in info:
                        info["work_years"] = next_cell_text
                elif "毕业学校" in cell_text and cell_text != "毕业学校":
                    if "school" not in info:
                        info["school"] = next_cell_text
                elif "专业" in cell_text and cell_text != "专业" and cell_text != "专业" * len(cell_text):
                    if "school" in info:
                        info["school"] += f" {next_cell_text}"
                    else:
                        info["major"] = next_cell_text
                elif "职称" in cell_text and cell_text != "职称":
                    if "title" not in info:
                        info["title"] = next_cell_text
                elif "毕业时间" in cell_text and cell_text != "毕业时间":
                    if "graduation_date" not in info:
                        info["graduation_date"] = next_cell_text
                elif "所在部门" in cell_text and cell_text != "所在部门":
                    if "department" not in info:
                        info["department"] = next_cell_text
                elif "个人简介" in cell_text and cell_text != "个人简介":
                    if "summary" not in info:
                        info["summary"] = next_cell_text
    return info

def extract_project_info(source_doc):
    projects = []
    in_project_section = False
    
    for table in source_doc.tables:
        for row in table.rows:
            row_text = " ".join([cell.text for cell in row.cells])
            if "项目经历" in row_text:
                in_project_section = True
                continue
            
            if in_project_section:
                cells = [cell.text.strip() for cell in row.cells]
                if len(cells) >= 4 and cells[0] and cells[0] not in ["开始时间", "项目经历"]:
                    project = {
                        "start_date": cells[0],
                        "end_date": cells[1] if len(cells) > 1 else "",
                        "project_name": cells[2] if len(cells) > 2 else "",
                        "role": cells[3] if len(cells) > 3 else "",
                        "description": cells[4] if len(cells) > 4 else ""
                    }
                    projects.append(project)
    return projects

def create_resume_from_template(template_path, source_resume_path, output_path):
    shutil.copy2(template_path, output_path)
    template = Document(output_path)
    source_doc = Document(source_resume_path)
    
    personal_info = extract_personal_info(source_doc)
    projects = extract_project_info(source_doc)
    
    for i, table in enumerate(template.tables):
        for row_idx, row in enumerate(table.rows):
            for cell_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                
                if "姓名" in cell_text and "name" in personal_info:
                    if personal_info["name"]:
                        cell.text = personal_info["name"]
                elif "毕业学校和专业" in cell_text and "school" in personal_info:
                    if personal_info["school"]:
                        cell.text = personal_info["school"]
                elif "工作年限" in cell_text and "work_years" in personal_info:
                    if personal_info["work_years"]:
                        cell.text = personal_info["work_years"]
                elif "主要经历专业技能" in cell_text and "summary" in personal_info:
                    if personal_info["summary"]:
                        cell.text = personal_info["summary"]
    
    template.save(output_path)
    print(f"已生成: {output_path.name}")

def process_resumes():
    OUTPUT_DIR.mkdir(exist_ok=True)
    
    for resume_file, original_file in RESUME_MAPPINGS.items():
        resume_path = RESUME_1_0_DIR / resume_file
        original_path = ORIGINAL_RESUMES_DIR / original_file
        
        if not resume_path.exists():
            print(f"警告: 简历文件不存在: {resume_file}")
            continue
            
        if not original_path.exists():
            print(f"警告: 原始简历不存在: {original_file}, 使用当前简历内容")
            original_path = resume_path
        
        output_path = OUTPUT_DIR / resume_file
        create_resume_from_template(TEMPLATE_PATH, original_path, output_path)

if __name__ == "__main__":
    process_resumes()
    print(f"\n所有简历处理完成！输出目录: {OUTPUT_DIR}")
